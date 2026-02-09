#!/usr/bin/env python3
"""
Amazon.co.jp ASIN Scraper
- Extracts product titles and bullet points from Amazon.co.jp
- Generates search keywords based on product data
- Searches Amazon.co.jp and collects results from 3 pages
- Outputs everything to an Excel file
"""

import requests
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import time
import random
import re
import json
import urllib.parse

# ============================================================
# Configuration
# ============================================================

PARENT_ASINS = {
    "フルーティーグラムティント": [
        "B0CG5WNCL5", "B0BY8PF9CT", "B0CG5Y7V9N", "B0CG5VPPYP", "B0CG5TXHNG",
        "B0BY8RL2NM", "B0BY8SFSR6", "B0CG5VV9LG", "B0CG5WJTP3", "B0BY8P4ZDR",
        "B0F5GMN738", "B0DMR5H4T1", "B0CTTD7PZQ", "B0CM2T7Y4S", "B0BY8R6CN1",
        "B0BY8QBKSR", "B0BY8NVDX5", "B0CG5X3GD2", "B0CG5X7RC6", "B0CG5VWPD3",
        "B0BY8TD35S", "B0CG5V3HJS", "B0BY8MNY4H", "B0F5GLSWSK", "B0CG5V7CJ1",
    ],
    "マキシグレイヤーティント": [
        "B0FN7B5TT8", "B0FN79P6ST", "B0FN7B16K7", "B0FN78BK5X", "B0FN79N8M5",
        "B0FN7B7QVF", "B0FN7BGR2F", "B0FN79QHCH", "B0FN79R9KG", "B0FN78Q958",
        "B0FN7BW2XQ", "B0FN79Y1QT", "B0FN7BFW4Y", "B0FN78QC4H", "B0FN78MW5R",
        "B0FN79VHM6", "B0FN7BFBTS", "B0FN79WXS9", "B0FN7C652K", "B0FN7B666T",
        "B0FN78H3F8",
    ],
    "ボンディンググロウリップスティック": [
        "B0CMTDKH9B", "B0CMTDNTQR", "B0CMTDYP83", "B0CG5VFWDK", "B0CG5TSFMW",
        "B0CG5VB34H", "B0CG5W85QL", "B0F9KFHZ6H", "B0CG5VQ19D", "B0CG5WQ1CC",
        "B0F9KHZPXX", "B0CG5V21MQ", "B0F9KZGWJS", "B0F9KG17GG", "B0CG5WLP51",
        "B0CMTBQ5RS", "B0CG5W38NV", "B0CMTDHN6S", "B0CMT9YM3G",
    ],
}

HEADERS_LIST = [
    {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,image/apng,*/*;q=0.8",
        "Accept-Language": "ja-JP,ja;q=0.9,en-US;q=0.8,en;q=0.7",
        "Accept-Encoding": "gzip, deflate, br",
        "Connection": "keep-alive",
        "Upgrade-Insecure-Requests": "1",
        "Sec-Fetch-Dest": "document",
        "Sec-Fetch-Mode": "navigate",
        "Sec-Fetch-Site": "none",
        "Sec-Fetch-User": "?1",
        "Cache-Control": "max-age=0",
    },
    {
        "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/121.0.0.0 Safari/537.36",
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,*/*;q=0.8",
        "Accept-Language": "ja,en-US;q=0.7,en;q=0.3",
        "Accept-Encoding": "gzip, deflate, br",
        "Connection": "keep-alive",
        "Upgrade-Insecure-Requests": "1",
    },
    {
        "User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/119.0.0.0 Safari/537.36",
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Accept-Language": "ja-JP,ja;q=0.8,en;q=0.6",
        "Accept-Encoding": "gzip, deflate, br",
        "Connection": "keep-alive",
        "Upgrade-Insecure-Requests": "1",
    },
]

OUTPUT_FILE = "/workspace/Amazon_ASIN_Analysis.xlsx"


def get_random_headers():
    return random.choice(HEADERS_LIST).copy()


def random_delay(min_sec=2, max_sec=5):
    delay = random.uniform(min_sec, max_sec)
    time.sleep(delay)


def fetch_product_page(asin, retries=3):
    """Fetch an Amazon.co.jp product page by ASIN."""
    url = f"https://www.amazon.co.jp/dp/{asin}"
    for attempt in range(retries):
        try:
            headers = get_random_headers()
            session = requests.Session()
            response = session.get(url, headers=headers, timeout=15)
            if response.status_code == 200:
                return response.text
            elif response.status_code == 503:
                print(f"  [WARN] 503 for {asin}, attempt {attempt+1}/{retries}")
                time.sleep(5 * (attempt + 1))
            else:
                print(f"  [WARN] Status {response.status_code} for {asin}")
                time.sleep(3 * (attempt + 1))
        except Exception as e:
            print(f"  [ERROR] Request failed for {asin}: {e}")
            time.sleep(3 * (attempt + 1))
    return None


def parse_product_data(html, asin):
    """Parse title and bullet points from Amazon product page HTML."""
    soup = BeautifulSoup(html, "lxml")
    
    # Extract title
    title = ""
    title_el = soup.find("span", {"id": "productTitle"})
    if title_el:
        title = title_el.get_text(strip=True)
    else:
        # Try alternative title selectors
        title_el = soup.find("h1", {"id": "title"})
        if title_el:
            title = title_el.get_text(strip=True)
        else:
            title_el = soup.select_one("#title_feature_div span")
            if title_el:
                title = title_el.get_text(strip=True)
    
    # Extract bullet points (箇条書き)
    bullets = []
    feature_bullets = soup.find("div", {"id": "feature-bullets"})
    if feature_bullets:
        items = feature_bullets.find_all("span", class_="a-list-item")
        for item in items:
            text = item.get_text(strip=True)
            if text and not text.startswith("›"):
                bullets.append(text)
    
    # Alternative: try #productFactsDesktopExpander
    if not bullets:
        facts = soup.find("div", {"id": "productFactsDesktopExpander"})
        if facts:
            items = facts.find_all("span", class_="a-list-item")
            for item in items:
                text = item.get_text(strip=True)
                if text:
                    bullets.append(text)
    
    # Alternative: try ul.a-unordered-list within feature-bullets area
    if not bullets:
        ul_elements = soup.select("#feature-bullets ul li span.a-list-item")
        for el in ul_elements:
            text = el.get_text(strip=True)
            if text:
                bullets.append(text)
    
    # If still empty, try broader approach
    if not bullets:
        about_section = soup.find("div", {"id": "productDescription"})
        if about_section:
            text = about_section.get_text(strip=True)
            if text:
                bullets.append(text)

    bullet_text = "\n".join(bullets) if bullets else ""
    
    return {
        "asin": asin,
        "title": title,
        "bullets": bullet_text,
    }


def generate_search_keywords(products, parent_name):
    """Generate 10 search keywords based on product data."""
    # Collect all titles and bullet points
    all_titles = [p["title"] for p in products if p["title"]]
    all_bullets = [p["bullets"] for p in products if p["bullets"]]
    
    # Extract common words/phrases from titles
    # Common patterns for cosmetics: brand, product type, color, features
    word_freq = {}
    
    for title in all_titles:
        # Clean and split
        # Remove color/shade specific parts (often in parentheses or after specific patterns)
        clean_title = re.sub(r'[（(][^）)]*[）)]', '', title)
        clean_title = re.sub(r'[0-9]+[gGmMlL]', '', clean_title)
        words = re.split(r'[\s　/／・]+', clean_title)
        for word in words:
            word = word.strip()
            if len(word) >= 2:
                word_freq[word] = word_freq.get(word, 0) + 1
    
    # Also parse bullet points for keywords
    bullet_words = {}
    for bullet in all_bullets:
        words = re.split(r'[\s　/／・\n]+', bullet)
        for word in words:
            word = word.strip()
            if len(word) >= 2:
                bullet_words[word] = bullet_words.get(word, 0) + 1
    
    # Sort by frequency
    sorted_title_words = sorted(word_freq.items(), key=lambda x: x[1], reverse=True)
    sorted_bullet_words = sorted(bullet_words.items(), key=lambda x: x[1], reverse=True)
    
    # Extract brand name (usually appears in most titles)
    brand = ""
    for word, count in sorted_title_words:
        if count >= len(all_titles) * 0.5:
            brand = word
            break
    
    # Try to extract product type from parent name
    product_type = parent_name
    
    # Extract key features from bullet points
    features = []
    feature_keywords = ["保湿", "ツヤ", "発色", "リップ", "ティント", "グロス", "マット", 
                        "潤い", "唇", "カラー", "色", "落ちない", "落ちにくい", "長持ち",
                        "グラデーション", "密着", "フルーティー", "リキッド", "クリーミー",
                        "ベルベット", "高発色", "透明感", "艶", "光沢", "グロウ", "ボンディング",
                        "スティック", "マキシ", "グレイヤー", "レイヤード"]
    
    for word, count in sorted_bullet_words:
        for kw in feature_keywords:
            if kw in word and word not in features:
                features.append(word)
    
    # Generate search keywords - combining different aspects
    keywords = []
    
    # 1. Parent name itself (most relevant)
    keywords.append(parent_name)
    
    # 2. Brand + product type
    if brand and brand != parent_name:
        keywords.append(f"{brand} {product_type}")
    
    # 3. Generic product type searches
    if "ティント" in parent_name:
        keywords.append("リップティント 韓国コスメ")
        keywords.append("リップティント 人気")
        keywords.append("リップティント 落ちない 発色")
        keywords.append("ティント リップ ツヤ")
        keywords.append("韓国 リップティント グラデーション")
    elif "リップスティック" in parent_name or "スティック" in parent_name:
        keywords.append("リップスティック 韓国コスメ")
        keywords.append("リップスティック ツヤ 保湿")
        keywords.append("グロウ リップスティック 人気")
        keywords.append("リップ スティック 発色")
        keywords.append("韓国 リップスティック グロウ")
    
    # 4. Feature-based keywords
    if "フルーティー" in parent_name or "フルーティー" in " ".join(all_titles):
        keywords.append("フルーティー リップ ティント")
    if "グラム" in parent_name:
        keywords.append("グラムティント リップ")
    if "マキシ" in parent_name:
        keywords.append("マキシ リップティント")
    if "グレイヤー" in parent_name:
        keywords.append("レイヤード ティント リップ")
    if "ボンディング" in parent_name:
        keywords.append("ボンディング リップ グロウ")
    if "グロウ" in parent_name:
        keywords.append("グロウ リップ ツヤ 韓国")
    
    # 5. Brand-specific if we found one
    if brand:
        keywords.append(f"{brand} リップ")
        keywords.append(f"{brand} 新作")
    
    # 6. General cosmetics keywords related to the product
    keywords.append("韓国コスメ リップ 人気 ランキング")
    keywords.append("リップ 高発色 保湿")
    
    # Remove duplicates while preserving order
    seen = set()
    unique_keywords = []
    for kw in keywords:
        if kw not in seen:
            seen.add(kw)
            unique_keywords.append(kw)
    
    # Return top 10
    return unique_keywords[:10]


def search_amazon(query, page=1, retries=3):
    """Search Amazon.co.jp and return results."""
    encoded_query = urllib.parse.quote(query)
    url = f"https://www.amazon.co.jp/s?k={encoded_query}&page={page}"
    
    for attempt in range(retries):
        try:
            headers = get_random_headers()
            headers["Referer"] = "https://www.amazon.co.jp/"
            session = requests.Session()
            response = session.get(url, headers=headers, timeout=15)
            if response.status_code == 200:
                return response.text
            elif response.status_code == 503:
                print(f"  [WARN] 503 for search page {page}, attempt {attempt+1}")
                time.sleep(5 * (attempt + 1))
            else:
                print(f"  [WARN] Status {response.status_code} for search page {page}")
                time.sleep(3 * (attempt + 1))
        except Exception as e:
            print(f"  [ERROR] Search failed: {e}")
            time.sleep(3 * (attempt + 1))
    return None


def parse_search_results(html):
    """Parse search results from Amazon.co.jp search page."""
    soup = BeautifulSoup(html, "lxml")
    results = []
    
    # Find search result items
    items = soup.select('div[data-asin][data-component-type="s-search-result"]')
    
    for item in items:
        asin = item.get("data-asin", "")
        if not asin:
            continue
        
        # Get title
        title = ""
        title_el = item.select_one("h2 a span")
        if title_el:
            title = title_el.get_text(strip=True)
        else:
            title_el = item.select_one("h2 span")
            if title_el:
                title = title_el.get_text(strip=True)
        
        if not title:
            continue
        
        link = f"https://www.amazon.co.jp/dp/{asin}"
        
        results.append({
            "asin": asin,
            "title": title,
            "link": link,
        })
    
    return results


def style_workbook(wb):
    """Apply formatting to the workbook."""
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font_white = Font(bold=True, size=11, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    for ws in wb.worksheets:
        # Style headers
        for cell in ws[1]:
            cell.font = header_font_white
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = thin_border
        
        # Auto-width columns (approximate)
        for col in ws.columns:
            max_length = 0
            col_letter = col[0].column_letter
            for cell in col:
                if cell.value:
                    # For Japanese text, each character is roughly 2 units wide
                    cell_len = len(str(cell.value))
                    if max_length < cell_len:
                        max_length = cell_len
                cell.border = thin_border
                cell.alignment = Alignment(vertical='top', wrap_text=True)
            
            adjusted_width = min(max_length + 2, 60)
            ws.column_dimensions[col_letter].width = adjusted_width


def main():
    print("=" * 60)
    print("Amazon.co.jp ASIN Analysis Tool")
    print("=" * 60)
    
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)
    
    for parent_name, asins in PARENT_ASINS.items():
        print(f"\n{'='*60}")
        print(f"Processing: {parent_name}")
        print(f"{'='*60}")
        
        # ============================================================
        # Step 1: Scrape product data
        # ============================================================
        print(f"\n[Step 1] Scraping {len(asins)} products...")
        products = []
        
        for i, asin in enumerate(asins):
            print(f"  [{i+1}/{len(asins)}] Fetching {asin}...")
            html = fetch_product_page(asin)
            if html:
                data = parse_product_data(html, asin)
                products.append(data)
                print(f"    Title: {data['title'][:60]}..." if data['title'] else "    Title: [Not found]")
                print(f"    Bullets: {len(data['bullets'].split(chr(10))) if data['bullets'] else 0} items")
            else:
                products.append({"asin": asin, "title": "[取得失敗]", "bullets": ""})
                print(f"    [FAILED] Could not fetch data")
            
            if i < len(asins) - 1:
                random_delay(2, 4)
        
        # ============================================================
        # Step 2 & 3: Create Raw Data sheet with search keywords
        # ============================================================
        print(f"\n[Step 2] Creating Raw Data sheet...")
        
        # Truncate sheet name if too long (Excel limit is 31 chars)
        raw_sheet_name = f"{parent_name} Raw Data"
        if len(raw_sheet_name) > 31:
            raw_sheet_name = raw_sheet_name[:31]
        
        ws_raw = wb.create_sheet(title=raw_sheet_name)
        
        # Headers
        ws_raw["A1"] = "Title"
        ws_raw["B1"] = "箇条書き (Bullet Points)"
        ws_raw["C1"] = "検索キーワード (Search Keywords)"
        
        # Generate search keywords
        print(f"\n[Step 3] Generating search keywords...")
        keywords = generate_search_keywords(products, parent_name)
        print(f"  Generated {len(keywords)} keywords:")
        for idx, kw in enumerate(keywords):
            print(f"    {idx+1}. {kw}")
        
        # Write product data
        for row_idx, product in enumerate(products, start=2):
            ws_raw.cell(row=row_idx, column=1, value=product["title"])
            ws_raw.cell(row=row_idx, column=2, value=product["bullets"])
            # Write keyword in C column for each row (up to 10)
            if row_idx - 2 < len(keywords):
                ws_raw.cell(row=row_idx, column=3, value=keywords[row_idx - 2])
        
        # ============================================================
        # Step 4 & 5: Search with top keyword and create Result sheet
        # ============================================================
        top_keyword = keywords[0] if keywords else parent_name
        print(f"\n[Step 4] Searching Amazon.co.jp with: '{top_keyword}'")
        
        all_search_results = []
        for page_num in range(1, 4):  # 3 pages
            print(f"  Fetching search page {page_num}...")
            search_html = search_amazon(top_keyword, page=page_num)
            if search_html:
                page_results = parse_search_results(search_html)
                print(f"    Found {len(page_results)} results on page {page_num}")
                all_search_results.extend(page_results)
            else:
                print(f"    [FAILED] Could not fetch search page {page_num}")
            
            if page_num < 3:
                random_delay(3, 6)
        
        # Remove duplicates by ASIN
        seen_asins = set()
        unique_results = []
        for r in all_search_results:
            if r["asin"] not in seen_asins:
                seen_asins.add(r["asin"])
                unique_results.append(r)
        
        print(f"\n[Step 5] Creating Result sheet with {len(unique_results)} results...")
        
        result_sheet_name = f"{parent_name} Result"
        if len(result_sheet_name) > 31:
            result_sheet_name = result_sheet_name[:31]
        
        ws_result = wb.create_sheet(title=result_sheet_name)
        
        # Headers
        ws_result["A1"] = "ASIN"
        ws_result["B1"] = "Title"
        ws_result["C1"] = "ASIN Link"
        
        # Write search results
        for row_idx, result in enumerate(unique_results, start=2):
            ws_result.cell(row=row_idx, column=1, value=result["asin"])
            ws_result.cell(row=row_idx, column=2, value=result["title"])
            ws_result.cell(row=row_idx, column=3, value=result["link"])
        
        print(f"  Done with {parent_name}!")
        
        # Save intermediate progress
        wb.save(OUTPUT_FILE)
        print(f"  [Saved] Intermediate progress saved to {OUTPUT_FILE}")
    
    # ============================================================
    # Final: Style and save
    # ============================================================
    print(f"\n{'='*60}")
    print("Applying styles and saving final file...")
    style_workbook(wb)
    wb.save(OUTPUT_FILE)
    print(f"\nDone! File saved to: {OUTPUT_FILE}")
    print(f"Sheets: {wb.sheetnames}")


if __name__ == "__main__":
    main()
