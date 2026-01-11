import re
import random
import time
from dataclasses import dataclass

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook


INPUT_TEXT = r"""
[B0DXKB2LN1]

B0FWRBCV52,B0FL7LB2TM,B0CCJ4JL3C,B0DLKPVXV3,B08V8RRJKG,B0B1JPZYMN,B085T6JWF6,B099NT57TY,B092CF6ZV7,B0CP5VRFZB,B0CMCL6F7Z,B0CJ2KY6RD,B0DDXV5KV4,B0CPLQPK34,B0DY79QRPX,B08ZMZQDJ8,B097M6PM3V,B0DNJS6T6R,B0FQTXS3DF,B0DXS34CDN,B0DB5NZDPS,B0FMK6FYL7,B0BPG9NWJG,B0DBTLG15D

B0B3MTMZCC,B0CSJXS1XW,B0CXPV462H,B084CZRCN8,B0DXKB2LN1,B084JVK6S2,B0F4CCYWLR,B084JT2738,B0F5Q8WCW3,B0F7PXZGPL,B08LG72KZD,B0F6N6WW99,B0DPM4NRRT,B0CS322SYQ,B08JV3KG45,B08LGCN5TH

B0DKXCP6BY,B0DTK9NS8C,B09G267W27,B0DXVVTD7N,B0FB7S541G,B0DXDTLLCD,B0FLP8X3YM,B0FP2GWYGX,B0BG7BP2XR,B010FOFSH0,B01N0ARFGP,B0CM4B43DZ,B08HJD4FBP,B0DYCKY68S,B0DXPG6DSR,B09ZY8T55Y,B0BFQ9RD5B,B0DJB99QMW,B0FST8KBX9,B0F5WDXKYD,B0DK76N358,B0F995H79C,B0CVQDYH9R,B0DMNFSJGR,B0DQCJ9T9N,B0DXPYR45S,B0DDK82BF4,B0CW36HJ6B,B0C1RS8TYR,B0FH8WXD1W,B0CX8GT3KS,B0CX5CVDKZ,B0DZ6R2TBP,B0D37B81SX

B0F371R464,B0CXPTLVNQ,B0D7HGRF2Q,B0F66FJRQP,B0D2XYTYY9,B0D2Y2138N,B0CXPWD4L4,B0D1VHQ34V,B0D7HQG3LC,B0F6MR7TGS,B0F6MV27Z2,B0F6MYCNYN,B0F6MW8PKB,B0F6N6WW99,B0F2SJJY5T,B0F2SSD3M3,B0F2SZ8R56,B0D2Y2HHXW,B0DHW266GR,B0D2Y4H97R,B0D2Y1ZN5M,B0DSP2WR9S,B0DSP1C94M,B0F4CG6CW7

B0D82F2CXR,B07JKJY97Q,B09YDCCJBJ,B0F2YXBBFD,B0DCJNDZD3,B0FKGR24MC,B0D2XYLS65,B08CXYHQXX,B09C7NHGQM,B074QPGCNB,B0711DTHY2,B09QCDHY4J,B0DB7WC3LL,B07PY1Y997,B0DCN86JV2,B0BWR3RKB9,B0F9NX6JYP,B08QZ2PB43

B0D7HQG3LC,B0DXKB2LN1,B0CTQCZG9S,B08LG72KZD,B0CXPTLVNQ,B084JVK6S2,B0F2SZ8R56,B0CXPV462H,B0CS322SYQ,B0DSP2WR9S,B0D2Y2138N,B0F371R464,B0D7HGRF2Q,B0DPM4NRRT

[B0CTQC8S75]

B0D82F2CXR,B07JKJY97Q,B09YDCCJBJ,B0F2YXBBFD,B0DCJNDZD3,B0FKGR24MC,B0D2XYLS65,B08CXYHQXX,B09C7NHGQM,B074QPGCNB,B0711DTHY2,B09QCDHY4J,B0DB7WC3LL,B07PY1Y997,B0DCN86JV2,B0BWR3RKB9,B0F9NX6JYP,B08QZ2PB43

B0D7HQG3LC,B0DXKB2LN1,B0CTQCZG9S,B08LG72KZD,B0CXPTLVNQ,B084JVK6S2,B0F2SZ8R56,B0CXPV462H,B0CS322SYQ,B0DSP2WR9S,B0D2Y2138N,B0F371R464,B0D7HGRF2Q,B0DPM4NRRT

B0793JVFXK,B0F5GK5SD1,B08LKDDRDF,B0D1BNR7Q2,B07XHLNNWH,B0BCD91Z4M,B0D9YDRN1Q,B0071FIJCI,B01N7DYFQU,B00FJRBY54,B07K1PWSB9,B00IJSWSA0,B0CWTNMXLZ,B01MTMMG67,B00LR92JCA,B07YFBJWNX,B0C89YQN5H,B09C2D3KD7,B0FTYYZWHF,B0CGXDVN78,B0B5G56XX1,B0BXS6XFVB,B0CWGG8FJ8,B0875P1NLQ,B0CHHJ8TB1,B07YPWBSJ7,B09BHSJQHB,B0B9LP6NF7,B0B8Z51S5F,B0BT8NBJBK

B0F6MV27Z2,B0F6MYCNYN,B0F6MW8PKB,B0CXPV462H,B0F6N2G6PT,B0F6N6WW99,B0F62V8VJ6,B0F5Q8WCW3,B0F2SJJY5T,B0F2SSD3M3,B0F2SZ8R56,B0D2Y2HHXW,B0DHW266GR,B0D2Y4H97R,B0D2Y1ZN5M,B0DSP2WR9S,B0DPM4NRRT,B0CTQCZG9S,B0F4CCYWLR,B0F4CG6CW7

B0CS322SYQ,B07GSQV3MX,B079KPRFXZ,B08WXBMBBC,B09R4TF1VG,B0F6N2G6PT,B0FJX6CTVX,B0F1CZ6Z5V,B0G3P5RNPW,B0DHCB6WLP,B07T9F4NS9,B0875P1NLQ,B0F9YG3PLJ,B0D6MSZBRG,B0FQ5FN6T7,B0BR27XM13,B0DCGBLJBJ,B0CXPLXL14,B0FZZT9P95,B091CBCPBY,B0FFSBCGVG,B0DHWS3SD3,B0D874XX9P,B0FXFHWTCN,B0FGHLNGCW,B077K3K2K3,B0FWJR6J5N,B0FGJ3HTF1,B0G923LP1S,B0967JKDHS,B0DHRKZGXR,B0FT3J3ZL1,B0CVRWCTCN,B0DCTRWSW4,B0F2GJSMF2,B0D3Q8GPBS,B0DCTSHC3W,B0D924J2S3,B0DMVTR8HX

[B0DR861CSK]

B0D82F2CXR,B07JKJY97Q,B09YDCCJBJ,B0F2YXBBFD,B0DCJNDZD3,B0FKGR24MC,B0D2XYLS65,B08CXYHQXX,B09C7NHGQM,B074QPGCNB,B0711DTHY2,B09QCDHY4J,B0DB7WC3LL,B07PY1Y997,B0DCN86JV2,B0BWR3RKB9,B0F9NX6JYP,B08QZ2PB43

B0D7HQG3LC,B0DXKB2LN1,B0CTQCZG9S,B08LG72KZD,B0CXPTLVNQ,B084JVK6S2,B0F2SZ8R56,B0CXPV462H,B0CS322SYQ,B0DSP2WR9S,B0D2Y2138N,B0F371R464,B0D7HGRF2Q,B0DPM4NRRT

B0793JVFXK,B0F5GK5SD1,B08LKDDRDF,B0D1BNR7Q2,B07XHLNNWH,B0BCD91Z4M,B0D9YDRN1Q,B0071FIJCI,B01N7DYFQU,B00FJRBY54,B07K1PWSB9,B00IJSWSA0,B0CWTNMXLZ,B01MTMMG67,B00LR92JCA,B07YFBJWNX,B0C89YQN5H,B09C2D3KD7,B0FTYYZWHF,B0CGXDVN78,B0B5G56XX1,B0BXS6XFVB,B0CWGG8FJ8,B0875P1NLQ,B0CHHJ8TB1,B07YPWBSJ7,B09BHSJQHB,B0B9LP6NF7,B0B8Z51S5F,B0BT8NBJBK

B0F6MV27Z2,B0F6MYCNYN,B0F6MW8PKB,B0CXPV462H,B0F6N2G6PT,B0F6N6WW99,B0F62V8VJ6,B0F5Q8WCW3,B0F2SJJY5T,B0F2SSD3M3,B0F2SZ8R56,B0D2Y2HHXW,B0DHW266GR,B0D2Y4H97R,B0D2Y1ZN5M,B0DSP2WR9S,B0DPM4NRRT,B0CTQCZG9S,B0F4CCYWLR,B0F4CG6CW7
"""


ASIN_RE = re.compile(r"^[A-Z0-9]{10}$")


def extract_section(text: str, section: str) -> str:
    marker = f"[{section}]"
    start = text.find(marker)
    if start < 0:
        raise ValueError(f"Section not found: {section}")
    start = start + len(marker)
    # Find next section marker
    next_start = text.find("[", start)
    if next_start < 0:
        return text[start:]
    return text[start:next_start]


def parse_asins(section_text: str) -> list[str]:
    raw = re.split(r"[\s,]+", section_text.strip())
    asins: list[str] = []
    seen: set[str] = set()
    for token in raw:
        t = token.strip().upper()
        if not t:
            continue
        if not ASIN_RE.match(t):
            continue
        if t in seen:
            continue
        seen.add(t)
        asins.append(t)
    return asins


def parse_input(text: str, sections: list[str]) -> dict[str, list[str]]:
    out: dict[str, list[str]] = {}
    for sec in sections:
        out[sec] = parse_asins(extract_section(text, sec))
    return out


def parse_title_from_jina(text: str) -> str | None:
    # r.jina.ai output usually contains a line like:
    # "Amazon.co.jp: <product title> : <category>"
    for line in text.splitlines():
        if not line.startswith("Amazon.co.jp:"):
            continue
        title = line[len("Amazon.co.jp:") :].strip()
        if " : " in title:
            title = title.rsplit(" : ", 1)[0].strip()
        return title or None
    return None


@dataclass(frozen=True)
class FetchResult:
    asin: str
    title: str | None
    ok: bool
    status: int | None
    error: str | None


def _looks_like_blocked(html_text: str) -> bool:
    lowered = html_text.lower()
    if "captchacharacters" in lowered:
        return True
    if "ロボット" in html_text:
        return True
    if "このページへのアクセスが拒否されました" in html_text:
        return True
    if "お客様のリクエストを処理" in html_text:
        return True
    return False


def _normalize_title(s: str) -> str:
    return " ".join(s.strip().split())


def fetch_title_direct(asin: str, session: requests.Session) -> FetchResult:
    # Mobile product page tends to be smaller & less likely to trigger blocks.
    url = f"https://www.amazon.co.jp/gp/aw/d/{asin}"
    try:
        r = session.get(url, timeout=20, allow_redirects=True)
        status = r.status_code
        if status != 200:
            return FetchResult(asin=asin, title=None, ok=False, status=status, error=f"HTTP {status}")

        if _looks_like_blocked(r.text):
            return FetchResult(asin=asin, title=None, ok=False, status=status, error="Blocked/CAPTCHA")

        soup = BeautifulSoup(r.text, "lxml")
        # Mobile page usually uses #title
        node = soup.select_one("#title") or soup.select_one("#productTitle")
        if node and node.get_text(strip=True):
            return FetchResult(
                asin=asin,
                title=_normalize_title(node.get_text(" ", strip=True)),
                ok=True,
                status=status,
                error=None,
            )

        # Fallback to <title>
        if soup.title and soup.title.string:
            t = _normalize_title(soup.title.string)
            # Remove trailing Amazon.co.jp parts if present
            t = re.sub(r"\s*-\s*Amazon\.co\.jp\s*$", "", t).strip()
            if t:
                return FetchResult(asin=asin, title=t, ok=True, status=status, error=None)

        return FetchResult(asin=asin, title=None, ok=False, status=status, error="Title not found in HTML")
    except Exception as e:  # noqa: BLE001
        return FetchResult(asin=asin, title=None, ok=False, status=None, error=str(e))


def fetch_title_with_retries(asin: str, session: requests.Session, retries: int = 4) -> FetchResult:
    # polite pacing + exponential-ish backoff when blocked
    last: FetchResult | None = None
    for attempt in range(retries):
        if attempt > 0:
            sleep_s = min(20, (2.5 * (2**attempt)) + random.random() * 2)
            time.sleep(sleep_s)
        res = fetch_title_direct(asin, session)
        last = res
        if res.ok:
            return res
        if res.error and "Blocked" in res.error:
            continue
        # non-blocked failures: don't retry too aggressively
        if attempt < retries - 1:
            time.sleep(1.0 + random.random())
    return last or FetchResult(asin=asin, title=None, ok=False, status=None, error="Unknown")


def ensure_sheets(wb, sheet_names: list[str]):
    # Delete all sheets not in target list, and recreate target sheets fresh.
    existing = {ws.title for ws in wb.worksheets}
    for ws in list(wb.worksheets):
        if ws.title not in sheet_names:
            wb.remove(ws)
    # Remove any existing target sheets to ensure clean rewrite
    for name in sheet_names:
        if name in existing:
            ws = wb[name]
            wb.remove(ws)
    # Create in desired order
    for name in sheet_names:
        wb.create_sheet(title=name)


def write_sheet(ws, asins: list[str], title_map: dict[str, str | None]):
    ws["A1"] = "ASIN"
    ws["B1"] = "Title"
    for i, asin in enumerate(asins, start=2):
        ws.cell(row=i, column=1, value=asin)
        ws.cell(row=i, column=2, value=title_map.get(asin) or "")


def main():
    src = "/workspace/Test123.xlsx"
    dst = "/workspace/Test123_filled.xlsx"
    sheet_names = ["B0DXKB2LN1", "B0CTQC8S75", "B0DR861CSK"]

    data = parse_input(INPUT_TEXT, sheet_names)
    all_asins: list[str] = []
    seen: set[str] = set()
    for name in sheet_names:
        for asin in data[name]:
            if asin in seen:
                continue
            seen.add(asin)
            all_asins.append(asin)

    title_map: dict[str, str | None] = {asin: None for asin in all_asins}

    session = requests.Session()
    session.headers.update(
        {
            "User-Agent": (
                "Mozilla/5.0 (Linux; Android 13; Pixel 7) AppleWebKit/537.36 "
                "(KHTML, like Gecko) Chrome/120 Mobile Safari/537.36"
            ),
            "Accept-Language": "ja-JP,ja;q=0.9,en-US;q=0.8,en;q=0.7",
        }
    )

    failures: list[FetchResult] = []
    for idx, asin in enumerate(all_asins, start=1):
        res = fetch_title_with_retries(asin, session, retries=3)
        title_map[asin] = res.title
        if not res.ok:
            failures.append(res)
        # pacing to reduce blocking
        time.sleep(0.35 + random.random() * 0.55)
        if idx % 50 == 0:
            time.sleep(2.5 + random.random() * 2)
        if idx % 10 == 0:
            ok_cnt = idx - len(failures)
            print(f"Progress: {idx}/{len(all_asins)} (ok={ok_cnt}, failed={len(failures)})", flush=True)

    wb = load_workbook(src)
    ensure_sheets(wb, sheet_names)

    for name in sheet_names:
        ws = wb[name]
        write_sheet(ws, data[name], title_map)

    # Put the first target sheet as active
    wb.active = 0
    wb.save(dst)

    total = len(all_asins)
    print(f"Saved: {dst}")
    print(f"Unique ASINs fetched: {total}, failures: {len(failures)}")
    for name in sheet_names:
        print(f"{name}: rows={len(data[name])}")
    if failures:
        print("Failed ASINs (up to 25):")
        for fr in failures[:25]:
            print(f"- {fr.asin}: {fr.error} (status={fr.status})")


if __name__ == "__main__":
    main()

