#!/usr/bin/env python3
"""
Create an Excel file identical to the input workbook, plus a new Sheet3 pivot.

Sheet3 layout:
- Rows (Y-axis): Record Month (YYYY-MM)
- Columns (X-axis): Merchant Customer ID(Merchant Name)
- Within each merchant column group (series):
  SP Spend, SB Spend, SD Spend, GMS, Total Ad Spend / GMS
"""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


EXPECTED_SHEET1 = "Sheet1"
EXPECTED_SHEET2 = "Sheet2"
OUTPUT_SHEET = "Sheet3"


@dataclass(frozen=True)
class Columns:
    merchant_id: str = "Merchant Customer ID"
    record_month: str = "Record Month"
    sb_spend: str = "SB Ad Spend(¥)(2025 )"
    sd_spend: str = "SD Ad Spend(¥)(2025 )"
    sp_spend: str = "SP Ad Spend(¥)(2025 )"
    gms: str = "Net Ordered GMS(¥)(2025 )"


def _to_int_str(x: Any) -> str:
    if x is None or pd.isna(x):
        return ""
    # Excel often reads large integers as float
    try:
        return str(int(float(x)))
    except Exception:
        return str(x).strip()


def _month_to_yyyy_mm(x: Any) -> str:
    if x is None or pd.isna(x):
        return ""
    try:
        s = str(int(float(x)))
    except Exception:
        s = str(x).strip()
    if len(s) == 6 and s.isdigit():
        return f"{s[:4]}-{s[4:]}"
    return s


def build_pivot(input_path: Path, cols: Columns) -> pd.DataFrame:
    df1 = pd.read_excel(input_path, sheet_name=EXPECTED_SHEET1, engine="openpyxl")
    df2 = pd.read_excel(input_path, sheet_name=EXPECTED_SHEET2, engine="openpyxl")

    required1 = [cols.merchant_id, cols.record_month, cols.sp_spend, cols.sb_spend, cols.sd_spend, cols.gms]
    missing1 = [c for c in required1 if c not in df1.columns]
    if missing1:
        raise ValueError(f"Missing columns in {EXPECTED_SHEET1}: {missing1}")

    required2 = [cols.merchant_id, "Merchant Name"]
    missing2 = [c for c in required2 if c not in df2.columns]
    if missing2:
        raise ValueError(f"Missing columns in {EXPECTED_SHEET2}: {missing2}")

    df1 = df1.copy()
    df2 = df2.copy()

    df1["merchant_id_str"] = df1[cols.merchant_id].map(_to_int_str)
    df2["merchant_id_str"] = df2[cols.merchant_id].map(_to_int_str)
    df2["merchant_name_str"] = df2["Merchant Name"].fillna("").astype(str).map(lambda s: "" if s.strip().lower() == "nan" else s.strip())

    id_to_name = (
        df2[df2["merchant_id_str"] != ""]
        .drop_duplicates(subset=["merchant_id_str"], keep="first")
        .set_index("merchant_id_str")["merchant_name_str"]
        .to_dict()
    )

    df1["merchant_name_str"] = df1["merchant_id_str"].map(lambda mid: id_to_name.get(mid, "")).fillna("")
    df1["merchant_label"] = df1.apply(
        lambda r: f"{r['merchant_id_str']}({r['merchant_name_str']})" if r["merchant_name_str"] else r["merchant_id_str"],
        axis=1,
    )

    df1["month_label"] = df1[cols.record_month].map(_month_to_yyyy_mm)
    df1 = df1[(df1["merchant_id_str"] != "") & (df1["month_label"] != "")]

    # Ensure numeric
    for c in [cols.sp_spend, cols.sb_spend, cols.sd_spend, cols.gms]:
        df1[c] = pd.to_numeric(df1[c], errors="coerce").fillna(0.0)

    agg = (
        df1.groupby(["month_label", "merchant_label"], as_index=False)[
            [cols.sp_spend, cols.sb_spend, cols.sd_spend, cols.gms]
        ]
        .sum()
        .rename(
            columns={
                cols.sp_spend: "sp_spend",
                cols.sb_spend: "sb_spend",
                cols.sd_spend: "sd_spend",
                cols.gms: "gms",
            }
        )
    )
    agg["ad_spend_vs_gms"] = (agg["sp_spend"] + agg["sb_spend"] + agg["sd_spend"]) / agg["gms"].replace({0.0: pd.NA})

    wide = agg.set_index(["month_label", "merchant_label"])[
        ["sp_spend", "sb_spend", "sd_spend", "gms", "ad_spend_vs_gms"]
    ].unstack("merchant_label")

    # Columns currently (metric, merchant); swap to (merchant, metric) for Excel grouping
    wide.columns = wide.columns.swaplevel(0, 1)
    wide = wide.sort_index(axis=1, level=0)

    # Order months ascending (YYYY-MM sorts lexicographically OK)
    wide = wide.sort_index()
    return wide


def write_sheet3(workbook_path: Path, output_path: Path, pivot: pd.DataFrame) -> None:
    wb = openpyxl.load_workbook(workbook_path)

    # Replace if exists
    if OUTPUT_SHEET in wb.sheetnames:
        del wb[OUTPUT_SHEET]
    ws = wb.create_sheet(OUTPUT_SHEET)

    header_fill = PatternFill("solid", fgColor="F2F2F2")
    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    right = Alignment(horizontal="right", vertical="center")

    metric_order = ["sp_spend", "sb_spend", "sd_spend", "gms", "ad_spend_vs_gms"]
    metric_labels = {
        "sp_spend": "SP Spend(¥)",
        "sb_spend": "SB Spend(¥)",
        "sd_spend": "SD Spend(¥)",
        "gms": "GMS(¥)",
        "ad_spend_vs_gms": "Total Ad Spend / GMS",
    }
    currency_format = "¥#,##0"
    percent_format = "0.00%"

    merchants = list(dict.fromkeys([m for (m, _metric) in pivot.columns]))
    metrics = metric_order

    # Row 1: merchant labels (merged)
    # Row 2: metric labels
    ws.cell(row=1, column=1, value="").alignment = center
    ws.cell(row=2, column=1, value="Record Month").alignment = center
    ws.cell(row=2, column=1).font = Font(bold=True)
    ws.cell(row=2, column=1).fill = header_fill
    ws.cell(row=2, column=1).border = border

    start_col = 2
    for i, merchant in enumerate(merchants):
        group_start = start_col + i * len(metrics)
        group_end = group_start + len(metrics) - 1

        ws.merge_cells(start_row=1, start_column=group_start, end_row=1, end_column=group_end)
        c = ws.cell(row=1, column=group_start, value=merchant)
        c.font = Font(bold=True)
        c.fill = header_fill
        c.alignment = center
        c.border = border

        for j, metric in enumerate(metrics):
            col = group_start + j
            h = ws.cell(row=2, column=col, value=metric_labels[metric])
            h.font = Font(bold=True)
            h.fill = header_fill
            h.alignment = center
            h.border = border

    # Data rows start at row 3
    for r_idx, month in enumerate(pivot.index.tolist(), start=3):
        mcell = ws.cell(row=r_idx, column=1, value=month)
        mcell.alignment = center
        mcell.border = border

        for i, merchant in enumerate(merchants):
            for j, metric in enumerate(metrics):
                col = start_col + i * len(metrics) + j
                val = pivot.get((merchant, metric), pd.Series(index=pivot.index)).get(month, pd.NA)
                if pd.isna(val):
                    v = None
                else:
                    v = float(val)
                cell = ws.cell(row=r_idx, column=col, value=v)
                cell.border = border
                cell.alignment = right
                if metric == "ad_spend_vs_gms":
                    cell.number_format = percent_format
                else:
                    cell.number_format = currency_format

    # Freeze header rows + first column
    ws.freeze_panes = "B3"

    # Reasonable column widths
    ws.column_dimensions["A"].width = 12
    for col in range(2, 2 + len(merchants) * len(metrics)):
        ws.column_dimensions[get_column_letter(col)].width = 16

    # Make header rows taller for readability
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 36

    wb.save(output_path)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", default="source.xlsx", help="Input .xlsx (downloaded Google Sheet)")
    ap.add_argument("--output", default="output_with_sheet3.xlsx", help="Output .xlsx with new Sheet3")
    args = ap.parse_args()

    input_path = Path(args.input).resolve()
    output_path = Path(args.output).resolve()

    cols = Columns()
    pivot = build_pivot(input_path, cols)
    write_sheet3(input_path, output_path, pivot)
    print(f"Wrote: {output_path}")


if __name__ == "__main__":
    main()

