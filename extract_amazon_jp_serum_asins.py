import re
import sys
import time
from dataclasses import dataclass
from typing import Iterable, List, Optional, Set, Tuple

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from playwright.sync_api import Browser, Page, Playwright, sync_playwright


AMAZON_JP = "https://www.amazon.co.jp"


EXCLUDE_ASINS_RAW = (
    "B0DCGDL4LZ,B0FKTSCQNZ,B0CLXNBVBL,B01MDTVZTZ,B0F8MVKMD6,B0DC61XNLX,B0DG5PCJTP,"
    "B09R49GF59,B0F6MRZ146,B0F2HMWFL8,B0D3TBDN1S,B09SX9DN4T,B0C537YZWL,B0CG9DDBCB,"
    "B0CVQRJQCT,B0D93K1RSP,B0BXWSHCWX,B0F9P2ZPCW,B0C7KDP1N2,B0DWKF7GRJ,B0FPWWLT8R,"
    "B0DB5SW66N,B0DKXCP6BY,B0DTK9NS8C,B09G267W27,B0DXVVTD7N,B0FB7S541G,B0DXDTLLCD,"
    "B0FLP8X3YM,B0FP2GWYGX,B0BG7BP2XR,B010FOFSH0,B01N0ARFGP,B0CM4B43DZ,B08HJD4FBP,"
    "B0DYCKY68S,B0DXPG6DSR,B09ZY8T55Y,B0BFQ9RD5B,B0DJB99QMW,B0FST8KBX9,B0F5WDXKYD,"
    "B0DK76N358,B0F995H79C,B0CVQDYH9R,B0DMNFSJGR,B0DQCJ9T9N,B0DXPYR45S,B0DDK82BF4,"
    "B0CW36HJ6B,B0C1RS8TYR,B0FH8WXD1W,B0CX8GT3KS,B0CX5CVDKZ,B0DZ6R2TBP,B0D37B81SX"
)


ASIN_RE = re.compile(r"^[A-Z0-9]{10}$")


def parse_exclude_asins(raw: str) -> Set[str]:
    items = [x.strip().upper() for x in raw.split(",") if x.strip()]
    bad = [x for x in items if not ASIN_RE.match(x)]
    if bad:
        raise ValueError(f"Invalid ASIN(s) in exclude list: {bad[:5]}{'...' if len(bad) > 5 else ''}")
    return set(items)


@dataclass(frozen=True)
class Item:
    asin: str
    title: str
    link: str


def _looks_like_robot_check(content: str) -> bool:
    markers = [
        "Enter the characters you see below",
        "ロボット",
        "captcha",
        "/errors/validateCaptcha",
        "Sorry, we just need to make sure you're not a robot",
    ]
    lc = content.lower()
    return any(m.lower() in lc for m in markers)


def _normalize_link(href: str) -> str:
    if not href:
        return ""
    if href.startswith("http://") or href.startswith("https://"):
        return href
    if href.startswith("/"):
        return f"{AMAZON_JP}{href}"
    return f"{AMAZON_JP}/{href}"


def _sleep_slightly(page_idx: int) -> None:
    # Keep it deterministic (no random), but not hammering.
    time.sleep(1.0 + 0.4 * page_idx)


def new_browser(playwright: Playwright) -> Tuple[Browser, Page]:
    browser = playwright.chromium.launch(
        headless=True,
        args=[
            "--disable-blink-features=AutomationControlled",
            "--no-sandbox",
        ],
    )
    context = browser.new_context(
        locale="ja-JP",
        viewport={"width": 1280, "height": 900},
        user_agent=(
            "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
            "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
        ),
    )
    # Speed up by blocking heavy resources.
    context.route(
        "**/*",
        lambda route: route.abort()
        if route.request.resource_type in {"image", "media", "font"}
        else route.continue_(),
    )
    page = context.new_page()
    page.set_default_timeout(45_000)
    return browser, page


def try_accept_consent(page: Page) -> None:
    # Japan site sometimes shows consent UI; best-effort.
    for label in ["同意する", "同意", "Accept", "OK"]:
        try:
            btn = page.get_by_role("button", name=label)
            if btn.count() > 0:
                btn.first.click(timeout=2_000)
                return
        except Exception:
            continue


def extract_items_from_html(html: str) -> List[Item]:
    soup = BeautifulSoup(html, "lxml")
    items: List[Item] = []

    for node in soup.select('div[data-component-type="s-search-result"][data-asin]'):
        asin = (node.get("data-asin") or "").strip().upper()
        if not ASIN_RE.match(asin):
            continue

        title = ""
        link = ""
        # Link + title can vary by locale/experiment.
        # 1) Preferred: h2 a span (common on many Amazon layouts)
        h2a = node.select_one("h2 a")
        if h2a is not None:
            link = _normalize_link(h2a.get("href") or "")
            title_span = h2a.select_one("span")
            if title_span is not None:
                title = title_span.get_text(strip=True)

        # 2) Fallback: first product-like anchor (often present without <h2>)
        if not link:
            for a in node.select("a.a-link-normal"):
                href = a.get("href") or ""
                if "/dp/" in href or "/gp/" in href:
                    link = _normalize_link(href)
                    break

        # 3) Fallback title: image alt (frequently includes the product title)
        if not title:
            img = node.select_one("img[alt]")
            if img is not None:
                title = (img.get("alt") or "").strip()

        if not link:
            link = f"{AMAZON_JP}/dp/{asin}"

        items.append(Item(asin=asin, title=title, link=link))

    return items


def scrape_serum_asins(max_pages: int = 5) -> List[Item]:
    query = "セラム"
    collected: List[Item] = []
    seen: Set[str] = set()

    with sync_playwright() as p:
        browser, page = new_browser(p)
        try:
            for page_idx in range(1, max_pages + 1):
                url = f"{AMAZON_JP}/s?k={query}&page={page_idx}"
                print(f"[{page_idx}/{max_pages}] Loading: {url}", flush=True)
                page.goto(url, wait_until="domcontentloaded", timeout=45_000)
                print(f"[{page_idx}/{max_pages}] DOMContentLoaded. Current URL: {page.url}", flush=True)
                try_accept_consent(page)

                html = page.content()
                print(f"[{page_idx}/{max_pages}] HTML size: {len(html)}", flush=True)
                if _looks_like_robot_check(html):
                    # Save debug artifacts and stop early.
                    page.screenshot(path=f"amazon_robot_check_page_{page_idx}.png", full_page=True)
                    with open(f"amazon_robot_check_page_{page_idx}.html", "w", encoding="utf-8") as f:
                        f.write(html)
                    raise RuntimeError(
                        "Amazon robot check detected. "
                        f"Saved amazon_robot_check_page_{page_idx}.png/.html for review."
                    )

                page_items = extract_items_from_html(html)
                print(f"[{page_idx}/{max_pages}] Extracted raw items: {len(page_items)}", flush=True)
                for it in page_items:
                    if it.asin in seen:
                        continue
                    seen.add(it.asin)
                    collected.append(it)
                print(f"[{page_idx}/{max_pages}] Unique so far: {len(collected)}", flush=True)

                _sleep_slightly(page_idx)
        finally:
            browser.close()

    return collected


def write_xlsx(items: Iterable[Item], out_path: str) -> None:
    wb = Workbook()
    ws: Worksheet = wb.active
    ws.title = "serum"

    ws["A1"] = "asin"
    ws["B1"] = "title"
    ws["C1"] = "link"

    row = 2
    for it in items:
        ws.cell(row=row, column=1, value=it.asin)
        ws.cell(row=row, column=2, value=it.title)
        ws.cell(row=row, column=3, value=it.link)
        row += 1

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 60

    wb.save(out_path)


def main(argv: List[str]) -> int:
    out_path = argv[1] if len(argv) > 1 else "amazon_jp_serum_asins.xlsx"
    max_pages = int(argv[2]) if len(argv) > 2 else 5

    exclude = parse_exclude_asins(EXCLUDE_ASINS_RAW)

    items = scrape_serum_asins(max_pages=max_pages)
    filtered = [it for it in items if it.asin not in exclude]

    write_xlsx(filtered, out_path=out_path)

    print(f"Scraped items (unique): {len(items)}")
    print(f"Excluded (provided list): {len(exclude)}")
    print(f"Written rows (after exclude): {len(filtered)}")
    print(f"Output: {out_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))

